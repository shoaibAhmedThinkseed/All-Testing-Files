import csv, argparse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

SPECIAL={(0x12,0x00),(0x22,0x01)}
HEADER=(0x4A,0x00)

def hx(v): return int(str(v).strip(),16)

def load_frames(path):
    frames=[]; cur=None
    with open(path,newline="",encoding="utf-8-sig") as f:
        r=csv.DictReader(f)
        if not r.fieldnames or "mosi" not in r.fieldnames or "miso" not in r.fieldnames:
            raise ValueError("CSV must contain mosi and miso columns")
        for row in r:
            t=row.get("type","").strip().lower()
            if t=="enable": cur=[]
            elif t=="result" and cur is not None: cur.append(row)
            elif t=="disable" and cur is not None:
                frames.append(cur); cur=None
    return frames

def parse(side,pos):
    out=[]; n=len(side); pending=0
    while pos<n:
        if all(x==0 for x in side[pos:]): break
        if pos+4>n: break
        tag=(side[pos],side[pos+1])
        outer=(side[pos+2]<<8)|side[pos+3]
        if tag in SPECIAL:
            # Keep tag, drop outer length and first 8 value bytes.
            out += [(side[pos],False),(side[pos+1],False)]
            vs=pos+4
            if vs+12>n: break
            # Keep connection ID.
            out += [(side[vs+8],False),(side[vs+9],False)]
            inner=(side[vs+10]<<8)|side[vs+11]
            ds=vs+12
            take=min(inner,n-ds)
            out += [(side[ds+i],True) for i in range(take)]
            pos=ds+take
            if take<inner:
                pending=inner-take; break
        else:
            # Preserve generic TLVs.
            out += [(side[pos+i],False) for i in range(4)]
            pos+=4
            take=min(outer,n-pos)
            out += [(side[pos+i],False) for i in range(take)]
            pos+=take
            if take<outer:
                pending=outer-take; break
    return out,pending

def consume_continuation(side,remaining):
    take=min(remaining,len(side))
    return [(side[i],True) for i in range(take)], remaining-take

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("input_csv"); ap.add_argument("output_xlsx")
    ap.add_argument("--header-len",type=int,default=19)
    a=ap.parse_args()

    frames=load_frames(a.input_csv)
    rows=[]; pending={"mosi":0,"miso":0}; frame=0
    phone_last=None; vehicle_last=None

    for tx in frames:
        mosi=[hx(x["mosi"]) for x in tx if x.get("mosi","").strip()]
        miso=[hx(x["miso"]) for x in tx if x.get("miso","").strip()]
        new_mosi=len(mosi)>=2 and tuple(mosi[:2])==HEADER
        new_miso=len(miso)>=2 and tuple(miso[:2])==HEADER

        if new_mosi or new_miso:
            frame+=1
            phone=[]; vehicle=[]
            if new_mosi:
                phone,pending["mosi"]=parse(mosi[a.header_len:],0)
            if new_miso:
                vehicle,pending["miso"]=parse(miso[a.header_len:],0)
            rows.append([frame,"BLE IC",phone])
            phone_last=len(rows)-1
            rows.append([frame,"Vehicle",vehicle])
            vehicle_last=len(rows)-1
        else:
            if pending["mosi"] and mosi:
                extra,pending["mosi"]=consume_continuation(mosi,pending["mosi"])
                if phone_last is not None: rows[phone_last][2].extend(extra)
            if pending["miso"] and miso:
                extra,pending["miso"]=consume_continuation(miso,pending["miso"])
                if vehicle_last is not None: rows[vehicle_last][2].extend(extra)

    wb=Workbook(); ws=wb.active; ws.title="Filtered SPI"
    maxn=max((len(x[2]) for x in rows),default=0)
    ws.append(["Frame","Direction"]+[f"Byte {i+1}" for i in range(maxn)])
    yellow=PatternFill("solid",fgColor="FFF176")
    for c in ws[1]:
        c.font=Font(bold=True); c.alignment=Alignment(horizontal="center")
    for frame,direction,data in rows:
        ws.append([frame,direction]+[""]*maxn); r=ws.max_row
        for i,(b,actual) in enumerate(data):
            c=ws.cell(r,i+3,f"{b:02X}"); c.alignment=Alignment(horizontal="center")
            if actual: c.font=Font(bold=True); c.fill=yellow
    ws.freeze_panes="C2"; ws.auto_filter.ref=ws.dimensions
    ws.column_dimensions["A"].width=10; ws.column_dimensions["B"].width=12
    for col in range(3,maxn+3):
        n=col; s=""
        while n: n,rem=divmod(n-1,26); s=chr(65+rem)+s
        ws.column_dimensions[s].width=8
    wb.save(a.output_xlsx)
    print("Created:",a.output_xlsx)

if __name__=="__main__": main()
